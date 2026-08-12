import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def create_report():
    doc = Document()

    # Page setup - Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("Jenkins Installation & Multi-Project Build Report")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("DevOps Practical Assignment | Stage-by-Stage Execution Document")
    sub_run.italic = True
    sub_run.font.size = Pt(12)
    sub_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Overview Box / Summary
    p = doc.add_paragraph()
    p.add_run("Overview: ").bold = True
    p.add_run("This document contains the step-by-step installation of Jenkins automation server and complete execution logs for three mandatory project types:\n")
    p.add_run("1. Freestyle Project using simple Windows Batch commands\n")
    p.add_run("2. Freestyle Project pulling source code from a GitHub repository\n")
    p.add_run("3. Maven Java Project executing automated build, test, and JAR packaging\n")

    img_dir = r"d:\SEM 7\Devops\screenshots"

    # Section 1: Installation
    h1 = doc.add_heading("1. Jenkins Installation and Controller Setup", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    p = doc.add_paragraph()
    p.add_run("Jenkins LTS (version 2.568.2) was downloaded and started as a standalone WAR service on ")
    p.add_run("http://localhost:8080").bold = True
    p.add_run(". The controller interface was verified after initial setup.")

    img1_path = os.path.join(img_dir, "01_jenkins_dashboard.png")
    if os.path.exists(img1_path):
        doc.add_paragraph("Figure 1: Jenkins Controller Dashboard UI").runs[0].bold = True
        doc.add_picture(img1_path, width=Inches(6.2))
        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 2: Freestyle Project
    h2 = doc.add_heading("2. Freestyle Project with Simple Commands", level=1)
    h2.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p = doc.add_paragraph()
    p.add_run("A new Freestyle project named ")
    p.add_run("Freestyle-Simple-Commands").bold = True
    p.add_run(" was created. A build step was configured using Windows Batch command syntax to execute basic system information commands.")

    code_p = doc.add_paragraph()
    code_run = code_p.add_run(
        'echo ==============================================\n'
        'echo Running Simple Command Freestyle Job\n'
        'echo Hostname: %COMPUTERNAME%\n'
        'echo Current Date and Time: %DATE% %TIME%\n'
        'echo ==============================================\n'
        'dir'
    )
    code_run.font.name = 'Consolas'
    code_run.font.size = Pt(9.5)

    img2_path = os.path.join(img_dir, "02_freestyle_project_creation.png")
    if os.path.exists(img2_path):
        doc.add_paragraph("Figure 2: Creating Freestyle Project").runs[0].bold = True
        doc.add_picture(img2_path, width=Inches(6.2))

    img3_path = os.path.join(img_dir, "03_freestyle_command_config.png")
    if os.path.exists(img3_path):
        doc.add_paragraph("Figure 3: Configuring Batch Build Steps").runs[0].bold = True
        doc.add_picture(img3_path, width=Inches(6.2))

    img4_path = os.path.join(img_dir, "04_freestyle_console_output.png")
    if os.path.exists(img4_path):
        doc.add_paragraph("Figure 4: Execution Console Output (SUCCESS)").runs[0].bold = True
        doc.add_picture(img4_path, width=Inches(6.2))
        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 3: GitHub Project
    h3 = doc.add_heading("3. Build Project from GitHub", level=1)
    h3.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p = doc.add_paragraph()
    p.add_run("A project named ")
    p.add_run("Build-From-GitHub").bold = True
    p.add_run(" was configured with Git Source Code Management. The repository URL ")
    p.add_run("https://github.com/jenkins-docs/simple-java-maven-app.git").bold = True
    p.add_run(" was cloned automatically during build execution.")

    img5_path = os.path.join(img_dir, "05_github_build_console_output.png")
    if os.path.exists(img5_path):
        doc.add_paragraph("Figure 5: Git Clone and SCM Checkout Console Output").runs[0].bold = True
        doc.add_picture(img5_path, width=Inches(6.2))
        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 4: Maven Project
    h4 = doc.add_heading("4. Build Maven Project", level=1)
    h4.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p = doc.add_paragraph()
    p.add_run("First, Apache Maven 3.9.9 was registered under ")
    p.add_run("Manage Jenkins -> Tools").bold = True
    p.add_run(". Then a project named ")
    p.add_run("Build-Maven-Project").bold = True
    p.add_run(" was created with top-level Maven targets ")
    p.add_run("clean package").bold = True
    p.add_run(". Jenkins successfully compiled the Java source files, executed JUnit unit tests, and generated the output JAR package.")

    img6_path = os.path.join(img_dir, "06_maven_global_tool_config.png")
    if os.path.exists(img6_path):
        doc.add_paragraph("Figure 6: Global Tool Configuration for Apache Maven 3.9.9").runs[0].bold = True
        doc.add_picture(img6_path, width=Inches(6.2))

    img7_path = os.path.join(img_dir, "07_maven_build_success_console.png")
    if os.path.exists(img7_path):
        doc.add_paragraph("Figure 7: Maven Build Success Console Output (JAR Generated)").runs[0].bold = True
        doc.add_picture(img7_path, width=Inches(6.2))
        doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 5: Summary Table
    h5 = doc.add_heading("5. Summary of Completed Deliverables", level=1)
    h5.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    headers = ["Task", "Job Name", "Source / Build Steps", "Status"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        set_cell_background(hdr_cells[i], '1F497D')
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data = [
        ("Freestyle Simple Commands", "Freestyle-Simple-Commands", "Batch Script (%COMPUTERNAME%, dir)", "Finished: SUCCESS"),
        ("Build from GitHub", "Build-From-GitHub", "Git SCM Checkout", "Finished: SUCCESS"),
        ("Build Maven Project", "Build-Maven-Project", "Maven Goals: clean package", "BUILD SUCCESS")
    ]

    for row_idx, row_data in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            if row_idx % 2 == 1:
                set_cell_background(row_cells[col_idx], 'F2F2F2')

    out_path = r"d:\SEM 7\Devops\Jenkins_Execution_Report.docx"
    doc.save(out_path)
    print(f"Report generated successfully at: {out_path}")

if __name__ == "__main__":
    create_report()
